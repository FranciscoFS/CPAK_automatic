"""Generate ABSTRACT_SCHOT_EN.docx with English figures."""
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# ── Title ──
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Development and Validation of a Deep Learning Algorithm for Automated Coronal Alignment Measurement in Long Limb X-rays")
run.bold = True
run.font.size = Pt(12)

# ── Authors ──
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Authors: Francisco Fernandez, Raúl Zilleruelo, Joaquín Steinsapir, Catalina Vidal, Cristian Ruz, Pablo Besa.")

# ── Sections ──
sections = [
    ("Introduction", "Coronal alignment assessment is a fundamental step in knee surgery planning, guiding decisions in corrective osteotomies and arthroplasty. It relies on measuring different angles on long limb X-rays (LLX), including the hip-knee-ankle angle (HKA), the mechanical lateral distal femoral angle (mLDFA), and the mechanical medial proximal tibial angle (mMPTA). These measurements characterize the coronal profile (varus, neutral, or valgus) and help phenotype patients for surgical planning. However, this process is time-consuming, requires training, and is subject to interobserver variability."),
    ("Objective", "To develop and validate a deep learning algorithm for automated coronal alignment measurement in LLX, compared against expert measurements."),
    ("Methods", "A deep learning algorithm based on computer vision models (YOLO) was developed, organized in three stages (Fig. 1): localization of the hip, knee, and ankle; identification of anatomical landmarks; and calculation of HKA, mLDFA, and mMPTA. The algorithm was trained with 280 radiographs for joint detection and 1,734 anatomical crops for landmark localization, using 80% for training and 20% for validation. For external validation, its measurements were compared with those of two experts in 112 limbs from 57 long limb X-rays, independent from training. Agreement was assessed using the intraclass correlation coefficient (ICC), specifically the two-way random-effects model for absolute agreement (ICC[2,1]), with 95% confidence intervals (95% CI). Bias was evaluated through Bland-Altman analysis."),
    ("Results", "The model achieved a mean average precision (mAP50) >97% for joint detection and >98% for landmark localization, with a processing time of 0.72 s per radiograph. Agreement between the model and expert measurements was excellent for HKA (ICC=0.995; 95% CI: 0.99\u20131.00) and good for mLDFA (ICC=0.873; 95% CI: 0.83\u20130.91) and mMPTA (ICC=0.873; 95% CI: 0.80\u20130.92). Interobserver agreement between the two experts was excellent for HKA (ICC=0.997; 95% CI: 0.99\u20131.00) and mLDFA (ICC=0.930; 95% CI: 0.87\u20130.96) and mMPTA (ICC=0.928; 95% CI: 0.89\u20130.95). Bland-Altman analysis showed biases of 0.20\u00b0, \u22120.15\u00b0, and \u22120.81\u00b0 for HKA, mLDFA, and mMPTA, respectively (Fig. 2). The mMPTA was the measurement with the largest bias and dispersion, consistent with its recognized difficulty of measurement even among expert surgeons."),
    ("Conclusion", "The model showed excellent performance and processing speed, with high agreement and low bias compared to expert measurements, approaching the level of interobserver agreement between the two experts. These findings suggest that AI-based models can standardize and optimize coronal alignment analysis in LLX."),
]

for title, body in sections:
    p = doc.add_paragraph()
    r = p.add_run(title + ": ")
    r.bold = True
    p.add_run(body)

# ── Figure 1 ──
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture("MANUSCRITO/Figure 1.png", width=Inches(6.0))
cap = doc.add_paragraph()
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = cap.add_run("Figure 1: Algorithm flowchart; from the input radiograph to angle calculation. HKA = Hip-Knee-Ankle angle; mLDFA = mechanical lateral distal femoral angle; mMPTA = mechanical medial proximal tibial angle.")
r.font.size = Pt(9)

# ── Figure 2 ──
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture("MANUSCRITO/Figure 2.png", width=Inches(6.0))
cap = doc.add_paragraph()
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = cap.add_run("Figure 2 Left column: Scatter plots of HKA, mLDFA, and mMPTA measurements from the model versus the observers' average. Right column: Bland-Altman plots for the same measurements.")
r.font.size = Pt(9)

doc.save("MANUSCRITO/ABSTRACT_SCHOT_EN.docx")
print("Guardado: MANUSCRITO/ABSTRACT_SCHOT_EN.docx")
